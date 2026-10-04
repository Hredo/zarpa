"""
Etapa 8 · Rasgos de las especies desde bases de datos revisadas por pares.

| Base | Grupo | Qué aporta | Licencia |
|---|---|---|---|
| AVONET (Tobias et al. 2022) | aves | nicho trófico, hábitat principal, modo de vida | CC BY 4.0 |
| EltonTraits 1.0 (Wilman et al. 2014) | aves y mamíferos | dieta en %, actividad (mamíferos), nocturnidad (aves) | sin restricciones |
| ReptTraits (Oskyrko et al. 2024) | reptiles | dieta, reproducción, actividad, microhábitat, hábitat | CC BY 4.0 |
| AmphiBIO (Oliveira et al. 2017) | anfibios | dieta, reproducción, actividad, hábitat | CC BY 4.0 |

Reglas:
  * Solo se asigna un rasgo si la base lo da para esa especie (nombre científico
    exacto, o el nombre aceptado en GBIF si iNaturalist usa otro). Sin dato, nada.
  * Reproducción de aves (todas ovíparas) y de mamíferos (vivíparos salvo los
    monotremas, ovíparos) es un hecho de toda la clase u orden: se asigna por
    clasificación y la procedencia lo dice.
  * Las categorías de dieta se derivan con umbrales escritos aquí (`_diet_*`),
    no por intuición, y la ficha muestra también el desglose de la fuente.

Salida: `out/traits.jsonl`.
"""

from __future__ import annotations

import csv
import json
from collections import defaultdict

from .. import config

T = config.EXTERNAL / "traits"

# Ambientes (máscara de bits). La app tiene la misma tabla en src/lib/groups.ts.
ENV = {
    "bosque": 1,
    "matorral": 2,
    "pradera": 4,
    "humedal": 8,
    "rios": 16,
    "costa": 32,
    "mar": 64,
    "roquedo": 128,
    "desierto": 256,
    "humanizado": 512,
    "ciudad": 1024,
    "granja": 2048,
    "selva": 4096,
}
MEDIUM = {"terrestre": 1, "agua_dulce": 2, "marino": 4, "salobre": 8}


def _num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# --- dieta -------------------------------------------------------------------

# Un recurso que da al menos el 60 % de la dieta define el nicho: es la regla de
# Pigot et al. 2020, la misma con la que AVONET asigna su nicho trófico. Si
# ninguno llega, se mira si la dieta es (casi) toda animal o toda vegetal.
SINGLE = [
    ("Inv", "Invertívoro"),
    ("Vfish", "Piscívoro"),
    ("Scav", "Carroñero"),
    ("Fruit", "Frugívoro"),
    ("Seed", "Granívoro"),
    ("Nect", "Nectarívoro"),
    ("PlantO", "Herbívoro"),
]


def _diet_from_percentages(p: dict[str, float]) -> tuple[str, str]:
    """EltonTraits: porcentajes por tipo de alimento → categoría + desglose."""
    vert = p["Vend"] + p["Vect"] + p["Vunk"]
    animal = p["Inv"] + vert + p["Vfish"] + p["Scav"]
    plant = p["Fruit"] + p["Nect"] + p["Seed"] + p["PlantO"]
    cat = next((label for k, label in SINGLE if p[k] >= 60), None)
    if cat is None and vert >= 60:
        cat = "Carnívoro"
    if cat is None:
        cat = "Carnívoro" if animal >= 90 else "Herbívoro" if plant >= 90 else "Omnívoro"
    labels = {
        "Inv": "invertebrados",
        "Vend": "aves y mamíferos",
        "Vect": "reptiles y anfibios",
        "Vfish": "peces",
        "Vunk": "otros vertebrados",
        "Scav": "carroña",
        "Fruit": "frutos",
        "Nect": "néctar",
        "Seed": "semillas",
        "PlantO": "otras partes de plantas",
    }
    parts = [f"{labels[k]} {int(v)} %" for k, v in sorted(p.items(), key=lambda kv: -kv[1]) if v >= 10]
    return cat, " · ".join(parts)


# Cuando dos fuentes nombran la dieta con distinto detalle pero sin
# contradecirse, se queda la más concreta.
COMPATIBLE = {("Depredador acuático", "Piscívoro"): "Piscívoro"}


def _agree(a: str | None, b: str | None) -> str | None:
    if a == b:
        return a
    return COMPATIBLE.get((a, b)) or COMPATIBLE.get((b, a))


AVONET_NICHE = {
    "Invertivore": "Invertívoro",
    "Vertivore": "Carnívoro",
    "Aquatic predator": "Depredador acuático",
    "Scavenger": "Carroñero",
    "Frugivore": "Frugívoro",
    "Granivore": "Granívoro",
    "Nectarivore": "Nectarívoro",
    "Herbivore terrestrial": "Herbívoro",
    "Herbivore aquatic": "Herbívoro",
    "Omnivore": "Omnívoro",
}

REPT_DIET = {"Carnivorous": "Carnívoro", "Herbivorous": "Herbívoro", "Omnivorous": "Omnívoro"}
ACTIVITY = {"Diurnal": "Diurno", "Nocturnal": "Nocturno", "Cathemeral": "De día y de noche", "Crepuscular": "Crepuscular"}
REPRO = {"oviparous": "Ovíparo", "viviparous": "Vivíparo", "ovoviviparous": "Ovovivíparo"}

AVONET_HABITAT = {
    "Forest": "bosque",
    "Woodland": "bosque",
    "Shrubland": "matorral",
    "Grassland": "pradera",
    "Wetland": "humedal",
    "Riverine": "rios",
    "Coastal": "costa",
    "Marine": "mar",
    "Rock": "roquedo",
    "Desert": "desierto",
    "Human Modified": "humanizado",
}
REPT_HABITAT = {
    "Forest": "bosque",
    "Woodlands": "bosque",
    "Shrubland": "matorral",
    "Savanna": "pradera",
    "Grassland": "pradera",
    "Wetlands": "humedal",
    "Rocky": "roquedo",
    "Desert": "desierto",
    "Marine Neritic": "mar",
    "Marine Oceanic": "mar",
    "Marine Intertidal": "costa",
    "Marine Coastal": "costa",
}
# Regiones biogeográficas de ReptTraits que son tropicales en (casi) toda su
# extensión. Australo-Pacífica y Saharo-Sindia mezclan climas: fuera.
TROPICAL_REALMS = {"Neotropic", "Oriental", "Afrotropic", "Madagascan"}

# EltonTraits, estrato de alimentación de los mamíferos: M = marino; suelo,
# trepador, arborícola y aéreo = vida en tierra firme.
MAM_STRATUM = {"M": "marino", "G": "terrestre", "S": "terrestre", "Ar": "terrestre", "A": "terrestre"}


def _avonet() -> dict[str, dict]:
    import openpyxl

    wb = openpyxl.load_workbook(T / "avonet.xlsx", read_only=True)
    ws = wb["AVONET1_BirdLife"]
    rows = ws.iter_rows(values_only=True)
    hdr = next(rows)
    i = {h: k for k, h in enumerate(hdr)}
    out = {}
    for r in rows:
        name = r[i["Species1"]]
        if not name:
            continue
        envs = 0
        hab = AVONET_HABITAT.get(r[i["Habitat"]])
        if hab:
            envs |= ENV[hab]
            # Selva: bosque como hábitat principal y área de distribución
            # centrada entre los trópicos (centroide de AVONET). Dos datos de la
            # misma fuente, no una suposición sobre la especie.
            lat = r[i["Centroid.Latitude"]]
            if hab == "bosque" and r[i["Habitat"]] == "Forest" and isinstance(lat, (int, float)) and abs(lat) <= 23.44:
                envs |= ENV["selva"]
        medium = 0
        if r[i["Habitat"]] == "Marine":
            medium |= MEDIUM["marino"]
        if r[i["Primary.Lifestyle"]] == "Aquatic" and r[i["Habitat"]] in ("Wetland", "Riverine"):
            medium |= MEDIUM["agua_dulce"]
        out[name] = {
            "diet": AVONET_NICHE.get(r[i["Trophic.Niche"]]),
            "envs": envs,
            "medium": medium,
            "migration": {1: "Sedentaria", 2: "Migradora parcial", 3: "Migradora"}.get(r[i["Migration"]]),
        }
    return out


def _elton(file: str, mammals: bool) -> dict[str, dict]:
    out = {}
    with open(T / file, encoding="latin-1") as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        for r in reader:
            name = (r.get("Scientific") or "").strip()
            if not name:
                continue
            p = {k: _num(r.get(f"Diet-{k}")) for k in ["Inv", "Vend", "Vect", "Vfish", "Vunk", "Scav", "Fruit", "Nect", "Seed", "PlantO"]}
            if sum(p.values()) < 90:  # sin dato de dieta
                continue
            cat, detail = _diet_from_percentages(p)
            entry = {"diet": cat, "diet_detail": detail}
            if mammals:
                stratum = MAM_STRATUM.get((r.get("ForStrat-Value") or "").strip())
                if stratum:
                    entry["medium"] = MEDIUM[stratum]
                n, c, d = (_num(r.get("Activity-Nocturnal")), _num(r.get("Activity-Crepuscular")), _num(r.get("Activity-Diurnal")))
                if d and not n:
                    entry["activity"] = "Diurno"
                elif n and not d:
                    entry["activity"] = "Nocturno" if not c else "Nocturno y crepuscular"
                elif n and d:
                    entry["activity"] = "De día y de noche"
                elif c:
                    entry["activity"] = "Crepuscular"
            else:
                if r.get("Nocturnal") == "1":
                    entry["activity"] = "Nocturno"
                if r.get("PelagicSpecialist") == "1":
                    entry["medium"] = MEDIUM["marino"]
            out[name] = entry
    return out


def _repttraits() -> dict[str, dict]:
    import openpyxl

    wb = openpyxl.load_workbook(T / "repttraits.xlsx", read_only=True)
    ws = wb["Data"]
    rows = ws.iter_rows(values_only=True)
    hdr = next(rows)
    i = {h: k for k, h in enumerate(hdr)}
    out = {}
    for r in rows:
        name = r[i["Species"]]
        if not name:
            continue
        envs = 0
        habitats = [p.strip() for p in str(r[i["Habitat type"]] or "").split("/")]
        for part in habitats:
            if part in REPT_HABITAT:
                envs |= ENV[REPT_HABITAT[part]]
        # Selva: bosque, región biogeográfica tropical y temperatura media anual
        # de su área de 18 °C o más (tres datos de ReptTraits; sin el último se
        # colarían los bosques templados de Chile o del sur de China).
        regions = {p.strip() for p in str(r[i["Main biogeographic region"]] or "").split("/")}
        mat = r[i["Mean Annual Temperature (°C)"]]
        if "Forest" in habitats and regions & TROPICAL_REALMS and isinstance(mat, (int, float)) and mat >= 18:
            envs |= ENV["selva"]
        micro = str(r[i["Microhabitat"]] or "").lower()
        medium = 0
        if any(m in micro for m in ("terrestrial", "arboreal", "fossorial", "saxicolous", "semi-aquatic", "semiaquatic")):
            medium |= MEDIUM["terrestre"]
        if "marine" in micro:
            medium |= MEDIUM["marino"]
        if "aquatic" in micro:  # incluye semiacuáticos
            medium |= MEDIUM["agua_dulce"]
        out[name] = {
            "diet": REPT_DIET.get(r[i["Diet "]]),
            "repro": REPRO.get(r[i["Reproductive mode"]]),
            "activity": ACTIVITY.get(r[i["Active time"]]),
            "envs": envs,
            "medium": medium,
        }
    return out


def _amphibio() -> dict[str, dict]:
    out = {}
    with open(T / "AmphiBIO_v1.csv", encoding="utf-8", errors="replace") as fh:
        for r in csv.DictReader(fh):
            name = (r.get("Species") or "").strip()
            if not name:
                continue
            flag = lambda k: r.get(k) == "1"  # noqa: E731
            animal = flag("Arthro") or flag("Vert")
            plant = flag("Leaves") or flag("Flowers") or flag("Seeds") or flag("Fruits")
            diet = None
            if animal and plant:
                diet = "Omnívoro"
            elif flag("Arthro") and not flag("Vert") and not plant:
                diet = "Invertívoro"
            elif animal:
                diet = "Carnívoro"
            elif plant:
                diet = "Herbívoro"
            repro = None
            if flag("Viv") and not (flag("Dir") or flag("Lar")):
                repro = "Vivíparo"
            elif (flag("Dir") or flag("Lar")) and not flag("Viv"):
                repro = "Ovíparo"
            acts = [a for a, k in (("Diurno", "Diu"), ("Nocturno", "Noc"), ("Crepuscular", "Crepu")) if flag(k)]
            activity = None
            if acts == ["Diurno"]:
                activity = "Diurno"
            elif acts == ["Nocturno"]:
                activity = "Nocturno"
            elif "Diurno" in acts and "Nocturno" in acts:
                activity = "De día y de noche"
            elif acts:
                activity = " y ".join(a.lower() for a in acts).capitalize()
            medium = 0
            if flag("Ter") or flag("Arb") or flag("Fos"):
                medium |= MEDIUM["terrestre"]
            if flag("Aqu"):
                medium |= MEDIUM["agua_dulce"]
            out[name] = {"diet": diet, "repro": repro, "activity": activity, "medium": medium}
    return out


def run() -> None:
    universe = [json.loads(l) for l in open(config.OUT / "universe.jsonl", encoding="utf-8")]
    taxa = {}
    for line in open(config.OUT / "taxa.jsonl", encoding="utf-8"):
        try:
            t = json.loads(line)
        except json.JSONDecodeError:
            continue
        taxa[t["inat_id"]] = t
    gbif_names = {}
    gbif_path = config.OUT / "gbif_taxa.jsonl"
    if gbif_path.exists():
        for line in open(gbif_path, encoding="utf-8"):
            g = json.loads(line)
            if g.get("gbif_name"):
                gbif_names[g["inat_id"]] = g["gbif_name"]

    print("[traits] leyendo bases…", flush=True)
    avonet = _avonet()
    elton_birds = _elton("BirdFuncDat.txt", mammals=False)
    elton_mammals = _elton("MamFuncDat.txt", mammals=True)
    rept = _repttraits()
    amph = _amphibio()
    print(f"[traits] AVONET {len(avonet)} · Elton aves {len(elton_birds)} · Elton mamíferos {len(elton_mammals)} · ReptTraits {len(rept)} · AmphiBIO {len(amph)}", flush=True)

    hits = defaultdict(int)
    conflicts: list[tuple[int, str, str]] = []
    out_path = config.OUT / "traits.jsonl"
    with open(out_path, "w", encoding="utf-8") as fh:
        for u in universe:
            t = taxa.get(u["inat_id"])
            if not t:
                continue
            anc = {a["rank"]: a["name"] for a in t["ancestors"]}
            cls, order = anc.get("class"), anc.get("order")
            names = [u["name"]] + ([gbif_names[u["inat_id"]]] if u["inat_id"] in gbif_names else [])
            rec: dict = {"inat_id": u["inat_id"]}
            src: dict[str, str] = {}

            def take(db: dict, code: str):
                for n in names:
                    if n in db:
                        return db[n], code
                return None, None

            def add_medium(bits: int, code: str) -> None:
                rec["medium"] = rec.get("medium", 0) | bits
                have = src.get("medium", "").split(",") if src.get("medium") else []
                if code not in have:
                    src["medium"] = ",".join(have + [code])

            if cls == "Aves":
                rec["repro"], src["repro"] = "Ovíparo", "clase"
                add_medium(MEDIUM["terrestre"], "clase")
                av, c = take(avonet, "avonet")
                if av:
                    for k in ("diet", "envs", "migration"):
                        if av.get(k):
                            rec[k] = av[k]
                            src[k] = c
                    if av.get("medium"):
                        add_medium(av["medium"], c)
                el, c = take(elton_birds, "eltontraits")
                if el:
                    rec["diet_detail"] = el.get("diet_detail")
                    src["diet_detail"] = c
                    if not rec.get("diet"):
                        rec["diet"], src["diet"] = el["diet"], c
                    else:
                        both = _agree(rec["diet"], el["diet"])
                        if both:
                            rec["diet"], src["diet"] = both, "avonet,eltontraits"
                        else:
                            conflicts.append((u["inat_id"], rec["diet"], el["diet"]))
                            rec.pop("diet")
                            src.pop("diet")
                    if el.get("activity"):
                        rec["activity"], src["activity"] = el["activity"], c
                    if el.get("medium"):
                        add_medium(el["medium"], c)
            elif cls == "Mammalia":
                if order == "Monotremata":
                    rec["repro"], src["repro"] = "Ovíparo", "orden"
                else:
                    rec["repro"], src["repro"] = "Vivíparo", "clase"
                el, c = take(elton_mammals, "eltontraits")
                if el:
                    for k in ("diet", "diet_detail", "activity"):
                        if el.get(k):
                            rec[k], src[k] = el[k], c
                    if el.get("medium"):
                        add_medium(el["medium"], c)
            elif cls == "Reptilia":
                rp, c = take(rept, "repttraits")
                if rp:
                    for k in ("diet", "repro", "activity", "envs"):
                        if rp.get(k):
                            rec[k], src[k] = rp[k], c
                    if rp.get("medium"):
                        add_medium(rp["medium"], c)
            elif cls == "Amphibia":
                am, c = take(amph, "amphibio")
                if am:
                    for k in ("diet", "repro", "activity"):
                        if am.get(k):
                            rec[k], src[k] = am[k], c
                    if am.get("medium"):
                        add_medium(am["medium"], c)
            if len(rec) > 1:
                rec["sources"] = src
                for k in src:
                    hits[k] += 1
                fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
    print(f"[traits] aves con dieta discrepante entre AVONET y EltonTraits (sin categoría): {len(conflicts)}")
    print(f"[traits] especies con rasgos por campo: {dict(hits)} → {out_path}")
